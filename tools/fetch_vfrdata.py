"""Données VFR de la carte de route -> vfr_data.json (embarqué dans l'appli : marche hors ligne).

- Points de report VFR (France + Belgique) : open flightmaps (OFMX, licence OFMA : utilisation
  libre, même commerciale, avec attribution « open flightmaps association »).
- Obstacles France : liste officielle du SIA (eAIP ENR 5.4, fichier ObstaclesDataZone1MFRANCE_*.xlsx,
  mise à jour à chaque cycle AIRAC), limitée à la zone de vol (nord de la France).
- Obstacles Belgique : eAIP skeyes ENR 5.4 (éoliennes à balisage militaire + fichier
  « additional info ») ; la base complète de l'IGN belge (eTOD) n'a pas de lien direct.
  OpenAIP (obstacles) en complément si l'API répond (elle est souvent saturée).
- Cartes VAC belges : liens vers les PDF officiels de l'eAIP skeyes (AD 2), par aérodrome.
  (Les VAC françaises suivent une adresse fixe du SIA, calculée par la page.)
- Cycle open flightmaps publié (pour les tuiles de la carte VFR).

Usage : python tools/fetch_vfrdata.py      (workflow update_aip.yml, chaque jeudi)
"""
import datetime as dt
import io
import json
import os
import re
import sys
import time
import zipfile
import xml.etree.ElementTree as ET

import requests

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "vfr_data.json")
ZONE = {"south": 48.0, "north": 52.2, "west": -0.5, "east": 7.0}   # Belgique + ~1 h 30 de vol autour
OFM = "https://snapshots.openflightmaps.org"
BE_EAIP = "https://ops.skeyes.be/html/belgocontrol_static/eaip/eAIP_Main/html/eAIP/"
SIA = "https://www.sia.aviation-civile.gouv.fr/media/dvd/"
UA = {"User-Agent": "Mozilla/5.0 (QuizPPL briefing; +https://str9ke.github.io/Quizz-PPL/)"}


def in_zone(lat, lon):
    return ZONE["south"] <= lat <= ZONE["north"] and ZONE["west"] <= lon <= ZONE["east"]


def ofm_cycle(s):
    """Cycle AIRAC le plus récent pour lequel open flightmaps a publié la Belgique."""
    now = dt.datetime.now(dt.timezone.utc)
    yy = now.year % 100
    for cand in [f"{yy}{n:02d}" for n in range(14, 0, -1)] + [f"{yy - 1}{n:02d}" for n in range(14, 10, -1)]:
        r = s.get(f"{OFM}/publicationServices/EBBU_{cand}.xml", timeout=30)
        if r.status_code != 200:
            continue
        for m in re.finditer(r'<cycle id="(\d+)"[^>]*startValidity="([^"]+)"[^>]*endValidity="([^"]+)"', r.text):
            a = dt.datetime.fromisoformat(m[2].replace("Z", "+00:00"))
            b = dt.datetime.fromisoformat(m[3].replace("Z", "+00:00"))
            if a <= now < b:
                return m[1]
    return None


def _deg(v):
    m = re.match(r"(\d+(?:\.\d+)?)([NSEW])", v or "")
    if not m:
        return None
    x = float(m[1])
    return -x if m[2] in "SW" else x


def vfr_points(s, cycle):
    pts, seen = [], set()
    for reg, short in (("ebbu", "eb"), ("lf", "lf")):
        url = f"{OFM}/live/{cycle}/ofmx/{reg}/latest/ofmx_{short}.zip"
        r = s.get(url, timeout=120)
        r.raise_for_status()
        z = zipfile.ZipFile(io.BytesIO(r.content))
        name = next(n for n in z.namelist() if n.endswith(".ofmx") and "/isolated/" in n)
        root = ET.fromstring(z.read(name))
        for d in root.iter("Dpn"):
            typ = (d.findtext("codeType") or "")
            if not typ.startswith("VFR"):
                continue
            uid = d.find("DpnUid")
            lat, lon = _deg(uid.findtext("geoLat")), _deg(uid.findtext("geoLong"))
            if lat is None or lon is None:
                continue
            ad = d.findtext("AhpUidAssoc/codeId") or ""
            ident = uid.findtext("codeId") or ""
            key = (ident, ad, round(lat, 4), round(lon, 4))
            if key in seen:
                continue
            seen.add(key)
            pts.append({"id": ident, "n": d.findtext("txtName") or ident, "ad": ad,
                        "c": 1 if typ == "VFR-MRP" else 0, "lat": round(lat, 5), "lon": round(lon, 5)})
    return pts


def airac_dir(today=None):
    today = today or dt.date.today()
    ref = dt.date(2026, 6, 11)
    d = ref + dt.timedelta(days=28 * ((today - ref).days // 28))
    mon = ["JAN", "FEB", "MAR", "APR", "MAY", "JUN", "JUL", "AUG", "SEP", "OCT", "NOV", "DEC"][d.month - 1]
    return f"{SIA}eAIP_{d.day:02d}_{mon}_{d.year}/FRANCE/", d


def obstacles_fr(s):
    import openpyxl
    base, d = airac_dir()
    page = s.get(f"{base}AIRAC-{d.isoformat()}/html/eAIP/FR-ENR-5.4-fr-FR.html", timeout=60).text
    m = re.search(r"(ObstaclesDataZone1MFRANCE_\d+\.xlsx)", page)
    if not m:
        raise RuntimeError("lien du fichier d'obstacles introuvable dans ENR 5.4")
    r = s.get(base + m[1], timeout=120)
    r.raise_for_status()
    wb = openpyxl.load_workbook(io.BytesIO(r.content), read_only=True)
    ws = wb.worksheets[0]
    rows = ws.iter_rows(values_only=True)
    head = [str(h or "").strip().upper() for h in next(rows)]
    col = {h: i for i, h in enumerate(head)}
    out = []
    for row in rows:
        try:
            lat = float(row[col["LATITUDE DECIMALE"]]); lon = float(row[col["LONGITUDE DECIMALE"]])
        except (TypeError, ValueError, KeyError):
            continue
        if not in_zone(lat, lon):
            continue
        unit = str(row[col["UNITE"]] or "FT").upper()
        k = 3.28084 if unit.startswith("M") else 1

        def ft(name):
            try:
                return int(round(float(row[col[name]]) * k))
            except (TypeError, ValueError, KeyError):
                return None
        lit = str(row[col.get("BALISAGE", -1)] or "") if "BALISAGE" in col else ""
        out.append({"id": str(row[col["IDENTIFICATEUR"]] or ""), "t": str(row[col["TYPE"]] or ""),
                    "nb": int(float(row[col["NOMBRE"]] or 1)), "lat": round(lat, 5), "lon": round(lon, 5),
                    "alt": ft("ALTITUDE AU SOMMET"), "h": ft("HAUTEUR HORS SOL"),
                    "lit": 0 if (not lit or "non" in lit.lower()) else 1, "src": "SIA"})
    return out, m[1]


def _dms(txt):
    """'501754N 0051036E' ou '50 34 06N', '003 38 26E' -> degrés."""
    t = re.sub(r"\s+", "", txt or "")
    m = re.match(r"(\d{2,3})(\d{2})(\d{2}(?:\.\d+)?)([NSEW])$", t)
    if not m:
        return None
    v = int(m[1]) + int(m[2]) / 60 + float(m[3]) / 3600
    return -v if m[4] in "SW" else v


def obstacles_be(get):
    import openpyxl
    out = []
    html = get("EB-ENR-5.4-en-GB.html")
    for m in re.finditer(r"<tr[^>]*>(.*?)</tr>", html, re.S):
        cells = [re.sub(r"<[^>]+>", " ", c) for c in re.findall(r"<td[^>]*>(.*?)</td>", m[1], re.S)]
        cells = [re.sub(r"\s+", " ", c).strip() for c in cells]
        if len(cells) < 4:
            continue
        pos = re.search(r"(\d{6}[NS])\s*(\d{7}[EW])", cells[2])
        hh = re.search(r"(\d+)\s*/\s*(\d+)", cells[3])
        if not pos or not hh:
            continue
        out.append({"id": "", "t": cells[1], "nb": 1, "lat": round(_dms(pos[1]), 5), "lon": round(_dms(pos[2]), 5),
                    "alt": int(hh[1]), "h": int(hh[2]), "lit": 1, "src": "skeyes", "rmk": cells[0] + " (balisage militaire)"})
    for f in sorted(set(re.findall(r"(https://ops\.skeyes\.be/[^\"'<> ]+additionalinfo\.xlsx)", html)))[-1:]:
        r = get(f, raw=True)
        if not r:
            continue
        wb = openpyxl.load_workbook(io.BytesIO(r), read_only=True)
        for row in wb.worksheets[0].iter_rows(values_only=True):
            if not row or str(row[0] or "").upper() != "AREA1":
                continue
            lat, lon = _dms(str(row[4] or "")), _dms(str(row[5] or ""))
            if lat is None or lon is None:
                continue
            out.append({"id": str(row[2] or ""), "t": "obstacle", "nb": 1, "lat": round(lat, 5), "lon": round(lon, 5),
                        "alt": row[8] if isinstance(row[8], (int, float)) else None,
                        "h": row[12] if isinstance(row[12], (int, float)) else None, "lit": 0, "src": "skeyes",
                        "rmk": "non vérifié par l'IGN"})
    return out


def obstacles_openaip(s):
    """Complément best-effort (Belgique) : l'API OpenAIP est souvent saturée -> on n'insiste pas."""
    key = os.environ.get("OPENAIP_API_KEY")
    if not key:
        return []
    out = []
    try:
        r = s.get("https://api.core.openaip.net/api/obstacles?bbox=2.5,49.45,6.45,51.55&limit=1000",
                  headers={"x-openaip-api-key": key}, timeout=60)
        if r.status_code != 200:
            print("OpenAIP obstacles : HTTP", r.status_code)
            return []
        for it in r.json().get("items", []):
            lon, lat = it["geometry"]["coordinates"][:2]
            h = (it.get("height") or {}).get("value")
            e = (it.get("elevation") or {}).get("value")
            out.append({"id": it.get("name", ""), "t": "obstacle", "nb": 1, "lat": round(lat, 5), "lon": round(lon, 5),
                        "alt": round(e * 3.28084) if e is not None and h is not None else None,
                        "h": round(h * 3.28084) if h is not None else None, "lit": 0, "src": "OpenAIP"})
    except Exception as e:
        print("OpenAIP obstacles :", e)
    return out


def vac_be(get):
    menu = get("EB-menu-en-GB.html")
    ads = sorted(set(re.findall(r"EB-AD-2\.([A-Z]{4})-en-GB\.html", menu)))
    res = {}
    for icao in ads:
        html = get(f"EB-AD-2.{icao}-en-GB.html")
        pdfs = []
        for p in sorted(set(re.findall(r"graphics/eAIP/([A-Z]{4}_(?:VAC|VLC|ADC|LDG)[^\"'<> ]*\.pdf)", html))):
            kind = re.search(r"_(VAC|VLC|ADC|LDG)", p)[1]
            pdfs.append({"k": kind, "u": "https://ops.skeyes.be/html/belgocontrol_static/eaip/eAIP_Main/graphics/eAIP/" + p})
        if not pdfs:
            # petits terrains sans carte dans l'AIP : site web de l'exploitant (AD 2.2)
            txt = re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", html))
            m = re.search(r"Website\s+((?:https?://)?(?:www\.)?[a-z0-9.-]+\.[a-z]{2,}[^\s]*)", txt, re.I)
            if m and m[1].upper() not in ("NIL", "INFO"):
                u = m[1] if m[1].lower().startswith("http") else "https://" + m[1]
                pdfs.append({"k": "WEB", "u": u})
        if pdfs:
            res[icao] = pdfs
        time.sleep(0.3)
    return res


def main():
    t0 = time.time()
    s = requests.Session()
    s.headers.update(UA)
    old = {}
    if os.path.exists(OUT):
        try:
            old = json.load(open(OUT, encoding="utf-8"))
        except ValueError:
            old = {}
    res = {"generated": dt.datetime.now(dt.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")}

    def keep(name, fn):
        try:
            res[name] = fn()
        except Exception as e:
            print(f"{name} en échec ({e}) : version précédente conservée")
            if name in old:
                res[name] = old[name]

    keep("ofm_airac", lambda: ofm_cycle(s))
    keep("vrp", lambda: vfr_points(s, res.get("ofm_airac")))

    fr_src = {}

    def _ofr():
        lst, src = obstacles_fr(s)
        fr_src["f"] = src
        return lst
    keep("obs_fr", _ofr)

    # eAIP skeyes : derrière le pare-feu anti-robots -> même session que fetch_notams.py
    try:
        sys.path.insert(0, ROOT)
        from fetch_notams import make_skeyes_session
        sk = make_skeyes_session()

        def get(page, raw=False):
            r = sk.get(page if page.startswith("http") else BE_EAIP + page, timeout=60)
            if r.status_code != 200:
                return b"" if raw else ""
            if raw:
                return r.content
            r.encoding = "utf-8"
            return r.text
        keep("obs_be", lambda: obstacles_be(get) + obstacles_openaip(s))
        keep("vac_be", lambda: vac_be(get))
    except Exception as e:
        print("session skeyes impossible :", e)
        for k in ("obs_be", "vac_be"):
            if k in old:
                res[k] = old[k]

    res["sources"] = {
        "vrp": "open flightmaps (OFMA General Users' License) — cycle " + str(res.get("ofm_airac")),
        "obs_fr": "SIA, eAIP ENR 5.4 — " + fr_src.get("f", (old.get("sources") or {}).get("obs_fr", "")),
        "obs_be": "skeyes eAIP ENR 5.4 (+ OpenAIP)",
        "vac_be": "skeyes eAIP AD 2",
    }
    tmp = OUT + ".tmp"
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(res, f, ensure_ascii=False, separators=(",", ":"))
    os.replace(tmp, OUT)
    print(f"vfr_data.json : cycle OFM {res.get('ofm_airac')} · {len(res.get('vrp', []))} points VFR · "
          f"{len(res.get('obs_fr', []))} obstacles FR · {len(res.get('obs_be', []))} obstacles BE · "
          f"VAC BE {len(res.get('vac_be', {}))} aérodromes — {time.time() - t0:.0f}s, "
          f"{os.path.getsize(OUT) // 1024} Ko")


if __name__ == "__main__":
    main()
